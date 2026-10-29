"""Export PowerPoint NATIF et MODIFIABLE (python-pptx).

Le mapping est reconstruit en formes PowerPoint (rectangles, ovales, zones de
texte) — pas une image collée : chaque texte, couleur et badge reste éditable
dans PowerPoint, Keynote ou LibreOffice. Géométrie identique à l'aperçu HTML
(1280×720 px ⇄ 13,333×7,5 pouces).
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

from config import theme
from core.models import Case
from render.html_renderer import MappingView, build_mapping_view

PX = 9525  # EMU par pixel CSS (96 dpi)
FONT = theme.PPTX_FONT


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper()[:6])


def _lighten(hex_color: str, f: float) -> str:
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    r, g, b = (round(c + (255 - c) * f) for c in (r, g, b))
    return f"#{r:02X}{g:02X}{b:02X}"


class Slide:
    def __init__(self, slide):
        self.s = slide

    def rect(self, x, y, w, h, fill: Optional[str], line: Optional[str] = None, radius: float = 0.0,
             shape=MSO_SHAPE.RECTANGLE, line_w: float = 0.75, name: str = ""):
        shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius and shape == MSO_SHAPE.RECTANGLE else shape
        sh = self.s.shapes.add_shape(shp_type, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
        if shp_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(line_w)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if name:
            sh.name = name
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return sh

    def text(self, x, y, w, h, paragraphs: Sequence, size: float = 10, color: str = theme.TEXT, bold=False,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, shape=None, line_spacing: float = 1.0, name: str = ""):
        """paragraphs : liste de str ou de listes de runs (texte, {size,bold,color})."""
        if shape is None:
            shape = self.s.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
        if name:
            shape.name = name
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(int(2 * PX))
        tf.margin_top = tf.margin_bottom = Emu(int(1 * PX))
        first = True
        for par in paragraphs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.line_spacing = line_spacing
            runs = par if isinstance(par, list) else [(par, {})]
            for txt, st in runs:
                r = p.add_run()
                r.text = str(txt)
                f = r.font
                f.name = FONT
                f.size = Pt(st.get("size", size))
                f.bold = st.get("bold", bold)
                f.color.rgb = _rgb(st.get("color", color))
        return shape

    def badge(self, x, y, w, h, text, fill, size=8.5, name=""):
        sh = self.rect(x, y, w, h, fill, radius=4, name=name)
        self.text(0, 0, 0, 0, [text], size=size, color="#FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE, shape=sh)
        return sh

    def line(self, x1, y1, x2, y2, color, width=1.5):
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * PX)), Emu(int(y1 * PX)),
                                         Emu(int(x2 * PX)), Emu(int(y2 * PX)))
        ln.line.color.rgb = _rgb(color)
        ln.line.width = Pt(width)
        return ln

    def oval(self, cx, cy, r, fill, line="#FFFFFF", line_w=2.0):
        return self.rect(cx - r, cy - r, 2 * r, 2 * r, fill, line, shape=MSO_SHAPE.OVAL, line_w=line_w)

    def pill(self, x, y, text):
        w = 18 + len(text) * 6.3
        self.badge(x, y, w, 20, text.upper(), theme.NAVY, 7.5)
        return w


def fit_pt(texts: Sequence[Tuple[str, float]], w_px: float, h_px: float, max_pt: float, min_pt: float = 6,
           lh: float = 1.22) -> float:
    """Taille (pt) la plus grande telle que les paragraphes tiennent dans la boîte.
    texts : (texte, facteur de taille relatif)."""
    size = max_pt
    while size > min_pt:
        total = 0.0
        for t, k in texts:
            px = size * k / 0.75
            cpl = max(1, int((w_px - 4) / (px * 0.53)))
            lines = 0
            for chunk in str(t).split("\n"):
                words, cur, n = chunk.split(), 0, 1
                for wd in words:
                    if cur and cur + 1 + len(wd) > cpl:
                        n += 1
                        cur = len(wd)
                    else:
                        cur += len(wd) + (1 if cur else 0)
                    n += max(0, math.ceil(len(wd) / cpl) - 1)
                lines += n
            total += lines * px * lh
        if total <= h_px - 2:
            return size
        size -= 0.25
    return min_pt


# ---------------------------------------------------------------------------
def add_mapping_slide(prs: Presentation, case: Case, view: Optional[MappingView] = None) -> None:
    v = view or build_mapping_view(case)
    S = Slide(prs.slides.add_slide(prs.slide_layouts[6]))
    two = len(v.rows) > 1
    W, H = 1280, 720
    S.rect(0, 0, W, H, theme.BG, name="Fond")

    # A. bandeau titre
    S.rect(0, 0, W, 56, theme.NAVY, name="Bandeau titre")
    tsize = fit_pt([(v.title, 1)], W - 40, 30, 17, 10)
    S.text(20, 6, W - 40, 30, [v.title], size=tsize, color="#FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE, name="Titre")
    S.text(20, 35, W - 40, 16, [v.subtitle], size=8, color="#D5DEE9", name="Sous-titre")

    # B. identité
    x0, y0, ww = 16, 65, W - 32
    S.rect(x0, y0, ww, 46, "#FFFFFF", theme.BORDER, radius=8, name="Identité")
    cw = ww / len(v.identity)
    for i, c in enumerate(v.identity):
        cx = x0 + i * cw
        if i:
            S.line(cx, y0 + 8, cx, y0 + 38, theme.BORDER, 0.75)
        S.text(cx + 2, y0 + 6, cw - 4, 14, [c["k"].upper()], size=6.5, color=theme.MUTED, bold=True)
        S.text(cx + 2, y0 + 19, cw - 4, 22, [c["v"]], size=fit_pt([(c["v"], 1)], cw - 6, 22, 11.5, 7),
               color=c.get("color") or theme.TEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # C. chronologie
    foot_y = H - 8 - 24
    ph = 226 if two else 258
    panels_y = foot_y - 8 - ph
    tl_top = 119
    S.pill(16, tl_top, "Chronologie du parcours")
    # légende
    S.rect(820, tl_top + 6, 16, 8, _lighten(theme.BLUE, .6))
    S.text(840, tl_top + 2, 150, 16, ["Fenêtre d'exposition probable"], size=7, color=theme.MUTED, align=PP_ALIGN.LEFT)
    S.rect(990, tl_top + 6, 16, 8, _lighten(theme.RED, .75))
    S.text(1010, tl_top + 2, 110, 16, ["Période contagieuse"], size=7, color=theme.MUTED, align=PP_ALIGN.LEFT)
    S.text(1110, tl_top + 2, 154, 16, ["J±n : jours depuis J0"], size=7, color=theme.MUTED, align=PP_ALIGN.RIGHT)

    rows_top = tl_top + 26
    rows_h = panels_y - 8 - rows_top
    gap_rows = 6 if two else 0
    row_h = (rows_h - gap_rows * (len(v.rows) - 1)) / max(1, len(v.rows))
    head_h = 34 if two else 46
    for ri, row in enumerate(v.rows):
        ry = rows_top + ri * (row_h + gap_rows)
        n = max(1, len(row))
        gap = 6
        cwid = (ww - gap * (n - 1)) / n
        band_y, card_y = ry, ry + 18
        card_h = row_h - 18
        for b in v.bands[ri]:
            bx = x0 + (b.col - 1) * (cwid + gap)
            bw = b.span * cwid + (b.span - 1) * gap
            col = _lighten(theme.BLUE, .78) if b.phase == "exposition" else _lighten(theme.RED, .85)
            tcol = theme.BLUE if b.phase == "exposition" else theme.RED
            sh = S.rect(bx, band_y, bw, 15, col, radius=4)
            S.text(0, 0, 0, 0, [b.label.upper()], size=6.5, color=tcol, bold=True, anchor=MSO_ANCHOR.MIDDLE, shape=sh)
        if row:
            dot_y = card_y + 4 + head_h + 12
            S.line(x0 + 10, dot_y, x0 + ww - 10, dot_y, "#C9D2DD", 1.5)
        # tailles communes (comme les groupes d'auto-ajustement HTML)
        body_h = card_h - 4 - head_h - 24 - (24 if two else 30) - 6
        head_pt = min(fit_pt([(c.date, 1), (c.place, .98)], cwid - 8, head_h - 4, 9.5 if not two else 8.5, 6) for c in row) if row else 8
        body_pt = min(fit_pt([(c.title, 1.02), (c.desc, .86)] + ([(" • ".join(ch["t"] for ch in c.chips), .74)] if c.chips else []),
                             cwid - 10, body_h, 10.5 if not two else 9, 5.5) for c in row) if row else 8
        for i, c in enumerate(row):
            cx = x0 + i * (cwid + gap)
            S.rect(cx, card_y, cwid, card_h, "#FFFFFF", theme.BORDER, radius=7, name=f"Étape {c.title}")
            S.rect(cx + 1, card_y + 0.5, cwid - 2, 4, c.type_color)
            S.text(cx + 2, card_y + 6, cwid - 4, head_h - 2,
                   [[(c.date, {"color": c.type_color, "bold": True, "size": head_pt})],
                    [(c.place, {"bold": True, "size": head_pt * .98})]], size=head_pt)
            dy = card_y + 4 + head_h + 12
            S.oval(cx + cwid / 2, dy, 10, c.type_color, "#FFFFFF", 2.5)
            if c.rel:
                rb = S.rect(cx + cwid - 30, dy - 8, 26, 12, theme.GREEN if c.rel == "J0" else theme.BG, radius=3)
                S.text(0, 0, 0, 0, [c.rel], size=6, bold=True, color="#FFFFFF" if c.rel == "J0" else theme.MUTED,
                       anchor=MSO_ANCHOR.MIDDLE, shape=rb)
            paras = [[(c.title, {"bold": True, "size": body_pt * 1.02})]]
            if c.desc:
                paras.append([(c.desc, {"color": theme.MUTED, "size": body_pt * .86})])
            if c.chips:
                paras.append([(" • ".join(ch["t"] for ch in c.chips), {"color": theme.NAVY, "size": body_pt * .74, "bold": True})])
            S.text(cx + 3, dy + 14, cwid - 6, body_h, paras, size=body_pt, line_spacing=1.0)
            bh = 18 if two else 22
            S.badge(cx + 6, card_y + card_h - bh - 6, cwid - 12, bh, c.risk, c.risk_color, 7.5 if two else 8)

    # D. panneaux
    fr = [1, 1.42, .86, .92]
    total_w = ww - 30
    widths = [total_w * f / sum(fr) for f in fr]
    xs = [x0]
    for wd in widths[:-1]:
        xs.append(xs[-1] + wd + 10)
    for x, wd, title in zip(xs, widths, ["Lieux & expositions", "Contacts à identifier / suivre",
                                          "PCI / décontamination", "Actions immédiates"]):
        S.rect(x, panels_y, wd, ph, "#FFFFFF", theme.BORDER, radius=8, name=title)
        S.pill(x + 11, panels_y + 9, title)
    inner_top = panels_y + 38
    inner_h = ph - 46

    # Lieux
    x, wd = xs[0], widths[0]
    n = max(1, len(v.places))
    rh = min(34, inner_h / n)
    pt = max(6, min(9.5, rh / 3.4))
    for i, p in enumerate(v.places):
        yy = inner_top + i * rh
        S.text(x + 8, yy, wd * .4, rh, [f'{p["n"]}. {p["name"]}'], size=pt, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        S.text(x + 8 + wd * .4, yy, wd * .32, rh, [p["popu"]], size=pt * .85, color=theme.MUTED, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        S.badge(x + wd - 8 - wd * .26, yy + rh * .18, wd * .26, rh * .64, p["risk"], p["color"], max(6, pt * .82))

    # Contacts
    x, wd = xs[1], widths[1]
    n_rows = max(1, math.ceil(len(v.contacts) / 2))
    rh = min(34, (inner_h - 34) / n_rows)
    pt = min(9.5, rh / 3.4)
    colw = (wd - 26) / 2
    for i, c in enumerate(v.contacts):
        col, r = i % 2, i // 2
        cx, yy = x + 8 + col * (colw + 10), inner_top + r * rh
        S.text(cx, yy, colw * .5, rh, [c["name"]], size=pt, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        num = [(c["n"], {"bold": True, "color": theme.RED if c["unknown"] else theme.NAVY, "size": pt * (1.1 if c["unknown"] else 1)})]
        if c["suffix"] and not c["unknown"]:
            num.append((" " + c["suffix"], {"color": theme.MUTED, "size": pt * .8}))
        S.text(cx + colw * .5, yy, colw * .23, rh, [num], size=pt, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        S.badge(cx + colw * .74, yy + rh * .18, colw * .26, rh * .64, c["risk"], c["color"], max(6, pt * .82))
    S.text(x + 8, inner_top + n_rows * rh + 4, wd - 16, 30, [v.contact_phrase], size=max(6.5, pt * .9), color=theme.RED, bold=True)

    # PCI
    x, wd = xs[2], widths[2]
    n = max(1, len(v.ipc))
    rh = min(32, (inner_h - 34) / n)
    pt = min(9.5, rh / 3.3)
    for i, s in enumerate(v.ipc):
        yy = inner_top + i * rh
        S.text(x + 8, yy, wd * .52, rh, [s["site"]], size=pt, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        S.badge(x + wd - 8 - wd * .4, yy + rh * .16, wd * .4, rh * .68, s["status"], s["color"], max(6, pt * .8))
    S.text(x + 8, inner_top + n * rh + 4, wd - 16, 34, [v.ipc_phrase], size=max(6.5, pt * .85), color=theme.MUTED, align=PP_ALIGN.LEFT)

    # Actions
    x, wd = xs[3], widths[3]
    n = max(1, len(v.actions))
    rh = min(28, (inner_h - (16 if v.actions_auto else 0)) / n)
    pt = fit_pt([(a["t"], 1) for a in v.actions], wd - 34, inner_h - (16 if v.actions_auto else 0), 10, 6, 1.6)
    for i, a in enumerate(v.actions):
        yy = inner_top + i * rh
        S.oval(x + 16, yy + rh / 2, 4.5, a["color"], a["color"], 0.5)
        S.text(x + 26, yy, wd - 32, rh, [a["t"]], size=pt, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
               color=theme.MUTED if a["done"] else theme.TEXT)
    if v.actions_auto:
        S.text(x + 10, panels_y + ph - 20, wd - 20, 14, ["Actions suggérées automatiquement — à valider."],
               size=6.5, color=theme.MUTED, align=PP_ALIGN.LEFT)

    # E. pied de page
    S.rect(x0, foot_y, ww, 24, "#E7EBF0", theme.BORDER, radius=5, name="Pied de page")
    S.text(x0 + 10, foot_y + 2, ww * .64, 20, [v.footer], size=7, color=theme.MUTED, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    S.text(x0 + ww * .64, foot_y + 2, ww * .36 - 10, 20, [v.meta], size=6.5, color=theme.MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(1280 * PX)
    prs.slide_height = Emu(720 * PX)
    return prs


def export_mapping_pptx(case: Case, path: Path | str) -> None:
    prs = new_presentation()
    add_mapping_slide(prs, case)
    prs.save(str(path))


def add_image_slide(prs: Presentation, png_path: Path | str) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(png_path), 0, 0, prs.slide_width, prs.slide_height)


def export_report_pptx(case: Case, images: List[Path], path: Path | str) -> None:
    """Rapport : mapping natif (modifiable) + visuels complémentaires en images HD."""
    prs = new_presentation()
    add_mapping_slide(prs, case)
    for img in images:
        add_image_slide(prs, img)
    prs.save(str(path))
