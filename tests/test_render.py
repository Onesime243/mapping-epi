import re
import shutil
import subprocess
from pathlib import Path
from xml.dom import minidom

import pytest
from pptx import Presentation

from core.models import Case, Event
from render import dashboard, gantt, html_renderer as H, map_view, network
from render.pptx_exporter import export_mapping_pptx


def test_mapping_html_structure(demo):
    html = H.render_mapping_html(demo)
    for s in ["CHRONOLOGIE", "Lieux &amp; expositions", "Contacts à identifier / suivre", "PCI / décontamination",
              "Actions immédiates", "3e CAS CONFIRMÉ MVE | ZS TUMBA", "Minimum documenté : 41 contacts"]:
        assert s.lower() in html.lower(), s
    assert len(re.findall(r'class="card[ "]', html)) == 12
    assert "Fenêtre d'exposition" in html


def test_title_and_two_rows_for_many_events():
    c = Case(rang=1, zs="Tumba", events=[Event(titre=f"Étape {i}") for i in range(15)])
    v = H.build_mapping_view(c)
    assert "DU 1er CAS" in v.title
    assert len(v.rows) == 2 and sum(len(r) for r in v.rows) == 15


def test_empty_case_renders():
    for fn in (H.render_mapping_html, network.render_html, map_view.render_html, gantt.render_html):
        assert "<svg" in fn(Case()) or "page mapping" in fn(Case())


def test_svgs_are_valid_xml(demo):
    for svg in (network.render_svg(demo), map_view.render_svg(demo), gantt.render_svg(demo),
                dashboard.render_svg([demo]), H.html_to_svg_wrapper(H.render_mapping_html(demo))):
        minidom.parseString(svg.encode("utf-8"))


def test_pptx_is_native_and_editable(demo, tmp_path):
    p = tmp_path / "m.pptx"
    export_mapping_pptx(demo, p)
    prs = Presentation(str(p))
    assert len(prs.slides) == 1
    shapes = prs.slides[0].shapes
    texts = " ".join(s.text_frame.text for s in shapes if s.has_text_frame)
    assert "Hospitalisation" in texts and "Minimum documenté : 41" in texts
    assert not any(s.shape_type == 13 for s in shapes)        # aucune image : 100 % formes éditables
    assert len(shapes) > 150


CHROME = next((p for p in [shutil.which("chromium"), shutil.which("google-chrome"), shutil.which("chrome"),
                           "/opt/pw-browsers/chromium_headless_shell-1243/chrome-headless-shell-linux64/chrome-headless-shell"]
               if p and Path(p).exists()), None)


@pytest.mark.skipif(CHROME is None, reason="navigateur headless indisponible")
def test_no_text_overflow(demo, tmp_path):
    """Aucun texte du mapping ne déborde après l'auto-ajustement (cas démo + cas surchargé)."""
    heavy = demo.model_copy(deep=True)
    for e in heavy.events:
        e.description = e.description + " • " + e.description
    for case in (demo, heavy):
        f = tmp_path / "m.html"
        f.write_text(H.render_mapping_html(case), encoding="utf-8")
        out = subprocess.run([CHROME, "--no-sandbox", "--headless", "--allow-file-access-from-files",
                              "--window-size=1280,720", "--virtual-time-budget=4000", "--dump-dom", f.as_uri()],
                             capture_output=True, text=True, timeout=60).stdout
        assert "__fitDone" in out or "page mapping" in out
        assert not re.search(r'class="[^"]*\boverflow\b', out), "texte tronqué dans le mapping"
