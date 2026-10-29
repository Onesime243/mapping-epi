"""Classeur Excel « modèle Mapping Épi » : saisie hors ligne puis import direct dans l'outil.

Structure (les noms de feuilles et d'en-têtes sont lus par l'import : ne pas les renommer) :
    Lisez-moi      mode d'emploi, règles de saisie, légende des couleurs
    Cas            identité, localisation, clinique (colonne « Valeur »)
    Chronologie    une ligne = une étape du parcours
    Lieux          lieux fréquentés et populations exposées
    Contacts       catégories de contacts (« ? » = nombre inconnu)
    PCI            sites et décontamination
    Laboratoire    prélèvements et délais
    Actions        actions immédiates
    Aperçu         image du mapping (facultative, ignorée à l'import)
    Listes / Ref_* listes déroulantes (masquées)

Dans chaque feuille-tableau : ligne 1 = titre, ligne 2 = consigne, ligne 3 = en-têtes, données dès la ligne 4.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import theme
from core.models import Case

HEADER_ROW = 3
FIRST_DATA_ROW = 4
MIN_ROWS = 40          # lignes vides prêtes à remplir dans chaque tableau

NAVY = theme.NAVY[1:]
WHITE = "FFFFFF"
MUTED = "5F6B7A"
SECTION_FILL = PatternFill("solid", fgColor="E3EDF9")
REQ_FILL = PatternFill("solid", fgColor="FFF4D6")
THIN = Side(style="thin", color="D3DAE3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RISK_FILLS = {
    "Très élevé": theme.RED, "Très haut": theme.RED, "Non": theme.RED, "Décès": theme.RED, "Urgent": theme.RED,
    "Élevé": theme.ORANGE, "Haut": theme.ORANGE, "Partiel": theme.ORANGE, "Important": theme.ORANGE,
    "Modéré": theme.GREEN, "Faible": theme.GREEN, "Contrôlé": theme.GREEN, "Décontaminé": theme.GREEN,
    "Inconnu": theme.BLUE, "À investiguer": theme.BLUE, "En cours": theme.BLUE, "Normal": theme.BLUE,
}

# --- Feuille Cas : (clé, libellé lu par l'import, aide, obligatoire, type) -----------------
CASE_SECTIONS = [
    ("IDENTIFICATION DU CAS", [
        ("rang", "Rang du cas", "Numéro d'ordre du cas confirmé dans la ZS (1, 2, 3…)", True, "int"),
        ("maladie", "Maladie", "Liste : MVE, Marburg, Mpox, Choléra, Rougeole, Autre", True, "maladie"),
        ("statut", "Statut", "Confirmé / Probable / Suspect", False, "statut"),
        ("sexe", "Sexe", "M / F / Inconnu", True, "sexe"),
        ("age", "Âge", "Nombre ; vide si inconnu", True, "int"),
        ("age_unite", "Unité âge", "ans / mois", False, "unite"),
        ("profession", "Profession", "", False, "text"),
    ]),
    ("LOCALISATION (pyramide sanitaire RDC)", [
        ("province", "Province", "Choisir dans la liste", True, "province"),
        ("zs", "Zone de santé", "Liste filtrée selon la province", True, "zs"),
        ("aire_sante", "Aire de santé", "Liste filtrée selon la ZS ; saisie libre si absente", True, "as"),
        ("village", "Village / quartier", "", False, "text"),
        ("lat", "Latitude", "GPS facultatif, ex. 2,6317", False, "float"),
        ("lon", "Longitude", "GPS facultatif, ex. 24,3912", False, "float"),
    ]),
    ("CLINIQUE ET ISSUE", [
        ("debut_symptomes", "Début symptômes", "J0 — JJ/MM/AAAA. Sert au calcul des fenêtres d'exposition", True, "date"),
        ("symptomes", "Symptômes (séparés par ;)", "ex. Fièvre; Céphalées; Méléna", False, "text"),
        ("symptomes_autres", "Autres symptômes", "", False, "text"),
        ("date_confirmation", "Date confirmation", "JJ/MM/AAAA", True, "date"),
        ("issue", "Issue", "Vivant / Décès / Guéri / En traitement / Inconnu", True, "issue"),
        ("date_deces", "Date décès", "JJ/MM/AAAA si décès", False, "date"),
        ("risque_global", "Risque global (vide = auto)", "Laisser vide : l'outil le calcule", False, "risque"),
    ]),
    ("INVESTIGATION", [
        ("investigateur", "Investigateur", "Nom / équipe", True, "text"),
        ("date_investigation", "Date investigation", "JJ/MM/AAAA", False, "date"),
    ]),
]

# --- Feuilles-tableaux : (feuille, titre, consigne, colonnes) ---------------------------------
# colonne = (clé, libellé lu par l'import, largeur, type, aide)
TABLES = [
    ("Chronologie", "CHRONOLOGIE DU PARCOURS",
     "Une ligne = une étape, dans l'ordre d'affichage. Date inconnue : laisser vide. Risque vide = calcul automatique.",
     [("date_debut", "Date début", 13, "date", "JJ/MM/AAAA — vide si inconnue"),
      ("date_fin", "Date fin", 13, "date", "Pour une période (ex. 01→14/08)"),
      ("heure", "Heure", 10, "text", "ex. 23h, 04h, 15–17h"),
      ("heure_approx", "Heure approx.", 9, "bool", "oui = affiché « ~23h »"),
      ("date_incertaine", "Date incertaine", 10, "bool", "oui = barre en pointillés"),
      ("lieu", "Lieu", 22, "text", "Lieu ou point de départ"),
      ("lieu_arrivee", "Lieu d'arrivée", 18, "text", "Pour un déplacement A → B"),
      ("type", "Type", 22, "event_type", "Détermine l'icône et la couleur"),
      ("titre", "Titre", 24, "text", "Texte en gras sur la carte"),
      ("description", "Description", 40, "text", "Séparer les éléments par •"),
      ("transport", "Transport", 12, "transport", "Moto, Véhicule, À pied…"),
      ("personnes", "Personnes présentes", 24, "text", "Population exposée à cette étape"),
      ("contact_corps", "Contact corps", 9, "bool", "oui = risque très élevé"),
      ("liquides", "Liquides biologiques", 10, "bool", "oui = +1 niveau de risque"),
      ("sans_pci", "Sans PCI", 9, "bool", "oui = +1 niveau de risque"),
      ("securise", "Sécurisé", 9, "bool", "EDS / transport sécurisé"),
      ("risque", "Risque (vide = auto)", 16, "risque", "Laisser vide pour le risque suggéré")]),
    ("Lieux", "LIEUX & EXPOSITIONS", "Lieux fréquentés par le cas et populations exposées.",
     [("nom", "Lieu", 26, "text", ""), ("type", "Type", 22, "place_type", ""),
      ("population", "Population exposée", 30, "text", "ex. Famille / communauté"),
      ("nombre", "Nombre estimé", 12, "int?", "Vide ou ? si inconnu"),
      ("risque", "Risque", 16, "risque", "Vide = calcul automatique"),
      ("lat", "Latitude", 12, "float", "Facultatif"), ("lon", "Longitude", 12, "float", "Facultatif")]),
    ("Contacts", "CONTACTS À IDENTIFIER / SUIVRE",
     "« ? » = nombre inconnu (jamais 0 par défaut). Le « Minimum documenté » est calculé par l'outil.",
     [("categorie", "Catégorie", 28, "text", "ex. Famille & amis, Motard(s)…"),
      ("lien", "Lien avec le cas", 20, "text", ""),
      ("nombre", "Nombre (? si inconnu)", 14, "int?", "Nombre ou ?"),
      ("libelle_nombre", "Libellé", 10, "text", "ex. listés"),
      ("suivis", "Suivis", 9, "int?", "Nombre suivi"),
      ("risque", "Risque", 12, "contact_risk", "Haut / Très haut"),
      ("lieu", "Lieu de contact", 22, "text", "Doit correspondre à un lieu de la feuille Lieux"),
      ("dernier_contact", "Dernier contact", 14, "date", "JJ/MM/AAAA — fin du suivi = +21 j")]),
    ("PCI", "PCI / DÉCONTAMINATION", "Documenter date, équipe et produit pour chaque site.",
     [("site", "Site", 28, "text", ""), ("statut", "Statut", 14, "ipc", "Décontaminé / Partiel / Non"),
      ("date", "Date", 13, "date", "JJ/MM/AAAA"), ("equipe", "Équipe", 22, "text", ""),
      ("produit", "Produit", 22, "text", "")]),
    ("Laboratoire", "LABORATOIRE", "Dates et heures au format JJ/MM/AAAA HH:MM. Le délai est calculé par l'outil.",
     [("type", "Échantillon", 14, "sample", ""), ("laboratoire", "Laboratoire", 20, "text", ""),
      ("prelevement", "Prélèvement", 17, "datetime", "JJ/MM/AAAA HH:MM"),
      ("envoi", "Envoi", 17, "datetime", ""), ("reception", "Réception", 17, "datetime", ""),
      ("resultat_le", "Résultat le", 17, "datetime", ""), ("resultat", "Résultat", 14, "result", "")]),
    ("Actions", "ACTIONS IMMÉDIATES", "Vide = l'outil propose des actions automatiquement.",
     [("texte", "Action", 46, "text", ""), ("priorite", "Priorité", 12, "priority", "Urgent / Important / Normal"),
      ("responsable", "Responsable", 22, "text", ""), ("echeance", "Échéance", 13, "date", "JJ/MM/AAAA"),
      ("statut", "Statut", 12, "action_status", "")]),
]
LIST_ATTR = {"Chronologie": "events", "Lieux": "places", "Contacts": "contacts", "PCI": "ipc",
             "Laboratoire": "labs", "Actions": "actions"}

LISTS = {
    "maladie": theme.DISEASES, "statut": theme.CASE_STATUS, "sexe": theme.SEXES, "unite": ["ans", "mois"],
    "issue": theme.ISSUE_CHOICES, "risque": theme.RISK_CHOICES, "event_type": theme.EVENT_TYPE_CHOICES,
    "transport": [t for t in theme.TRANSPORTS if t], "place_type": theme.PLACE_TYPES,
    "contact_risk": theme.CONTACT_RISK_CHOICES, "ipc": theme.IPC_CHOICES,
    "sample": ["Swab", "Sang total", "Plasma", "Urine", "Autre"],
    "result": ["Positif", "Négatif", "En attente", "Indéterminé"],
    "priority": theme.ACTION_PRIORITY_CHOICES, "action_status": theme.ACTION_STATUS_CHOICES, "bool": ["oui", "non"],
}


def _cell_value(v, kind: str):
    if isinstance(v, bool):
        return "oui" if v else None
    if isinstance(v, list):
        return "; ".join(v)
    return v


def _style_title(ws, text: str, sub: str, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(bold=True, size=15, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(row=2, column=1, value=sub)
    s.font = Font(italic=True, size=10, color=MUTED)
    s.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 22
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)


def _add_list_validation(ws, rng: str, list_name: str, list_ranges: dict, strict: bool = False):
    dv = DataValidation(type="list", formula1=list_ranges[list_name], allow_blank=True)
    dv.showErrorMessage = strict
    dv.error = "Choisissez une valeur dans la liste."
    ws.add_data_validation(dv)
    dv.add(rng)


def _risk_formatting(ws, rng: str):
    for value, color in RISK_FILLS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'],
                                                      fill=PatternFill("solid", fgColor=color[1:]),
                                                      font=Font(bold=True, color=WHITE)))


def build_workbook(case: Optional[Case] = None, preview_png: Optional[Path] = None) -> Workbook:
    from core import pyramide
    case = case or Case()
    wb = Workbook()

    # ---- listes (feuille masquée) --------------------------------------------------------
    lists = wb.active
    lists.title = "Listes"
    list_ranges = {}
    for col, (name, values) in enumerate(LISTS.items(), start=1):
        lists.cell(row=1, column=col, value=name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col, value=v)
        L = get_column_letter(col)
        list_ranges[name] = f"Listes!${L}$2:${L}${len(values) + 1}"
    pyr = pyramide.get()
    provs = pyr.provinces()
    pc = len(LISTS) + 1
    PL = get_column_letter(pc)
    lists.cell(row=1, column=pc, value="province").font = Font(bold=True)
    for r, p in enumerate(provs, start=2):
        lists.cell(row=r, column=pc, value=p)
    list_ranges["province"] = f"Listes!${PL}$2:${PL}${len(provs) + 1}"

    ref_zs = wb.create_sheet("Ref_ZS")
    ref_zs.append(["province", "zone_sante"])
    for p in provs:
        for z in pyr.zones(p):
            ref_zs.append([p, z])
    ref_as = wb.create_sheet("Ref_AS")
    ref_as.append(["province|zone", "aire_sante"])
    for p in provs:
        for z in pyr.zones(p):
            for a in pyr.aires(p, z):
                ref_as.append([f"{p}|{z}", a])

    # ---- Lisez-moi ---------------------------------------------------------------------
    readme = wb.create_sheet("Lisez-moi", 0)
    _readme(readme)

    # ---- Cas ---------------------------------------------------------------------------
    ws = wb.create_sheet("Cas", 1)
    _style_title(ws, "FICHE DU CAS — MAPPING DU PARCOURS ÉPIDÉMIOLOGIQUE",
                 "Remplir la colonne « Valeur ». Cases jaunes = indispensables. Inconnu : laisser vide (affiché « ? »).", 3)
    for i, h in enumerate(["Champ", "Valeur", "Aide"], start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        c.font, c.fill = Font(bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[HEADER_ROW].height = 22
    row = FIRST_DATA_ROW
    cell_of = {}
    for section, fields in CASE_SECTIONS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        c = ws.cell(row=row, column=1, value=section)
        c.font, c.fill = Font(bold=True, color=NAVY), SECTION_FILL
        c.alignment = Alignment(indent=1)
        row += 1
        for key, label, help_, required, kind in fields:
            a = ws.cell(row=row, column=1, value=label)
            a.font = Font(bold=required)
            a.alignment = Alignment(indent=1, vertical="center")
            v = ws.cell(row=row, column=2, value=_cell_value(getattr(case, key), kind))
            v.border = BOX
            v.alignment = Alignment(vertical="center", wrap_text=True)
            if required:
                v.fill = REQ_FILL
            if kind == "date":
                v.number_format = "DD/MM/YYYY"
            h = ws.cell(row=row, column=3, value=help_)
            h.font = Font(italic=True, size=9, color=MUTED)
            h.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row].height = 20
            cell_of[key] = f"$B${row}"
            if kind in LISTS:
                _add_list_validation(ws, f"B{row}", kind, list_ranges)
            elif kind == "province":
                _add_list_validation(ws, f"B{row}", "province", list_ranges)
            elif kind == "date":
                dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)",
                                    formula2="DATE(2100,12,31)", allow_blank=True)
                dv.error, dv.errorTitle = "Date au format JJ/MM/AAAA", "Date invalide"
                ws.add_data_validation(dv)
                dv.add(f"B{row}")
            row += 1
    # listes en cascade Province → ZS → AS
    prov, zs = cell_of["province"], cell_of["zs"]
    n_zs, n_as = ref_zs.max_row, ref_as.max_row
    f_zs = (f"OFFSET(Ref_ZS!$B$1,MATCH({prov},Ref_ZS!$A$2:$A${n_zs},0),0,"
            f"MAX(1,COUNTIF(Ref_ZS!$A$2:$A${n_zs},{prov})),1)")
    f_as = (f"OFFSET(Ref_AS!$B$1,MATCH({prov}&\"|\"&{zs},Ref_AS!$A$2:$A${n_as},0),0,"
            f"MAX(1,COUNTIF(Ref_AS!$A$2:$A${n_as},{prov}&\"|\"&{zs})),1)")
    for key, formula in (("zs", f_zs), ("aire_sante", f_as)):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.showErrorMessage = False          # une aire absente de la liste reste saisissable
        ws.add_data_validation(dv)
        dv.add(cell_of[key].replace("$", ""))
    _risk_formatting(ws, cell_of["risque_global"].replace("$", ""))
    _risk_formatting(ws, cell_of["issue"].replace("$", ""))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 58
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.sheet_view.showGridLines = False
    _print_setup(ws, HEADER_ROW, landscape=False)

    # ---- feuilles-tableaux -------------------------------------------------------------
    for idx, (sheet, title, hint, cols) in enumerate(TABLES, start=2):
        sh = wb.create_sheet(sheet, idx)
        _style_title(sh, title, hint, len(cols))
        for i, (key, label, width, kind, tip) in enumerate(cols, start=1):
            c = sh.cell(row=HEADER_ROW, column=i, value=label)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            if tip:
                c.comment = Comment(tip, "Mapping Épi", width=220, height=60)
            sh.column_dimensions[get_column_letter(i)].width = width
        sh.row_dimensions[HEADER_ROW].height = 34
        items = getattr(case, LIST_ATTR[sheet])
        for r, item in enumerate(items, start=FIRST_DATA_ROW):
            for i, (key, label, width, kind, tip) in enumerate(cols, start=1):
                v = getattr(item, key)
                if kind == "int?" and v is None and sheet == "Contacts" and key == "nombre":
                    v = "?"
                sh.cell(row=r, column=i, value=_cell_value(v, kind))
        last = FIRST_DATA_ROW + max(len(items), MIN_ROWS) - 1
        for i, (key, label, width, kind, tip) in enumerate(cols, start=1):
            L = get_column_letter(i)
            rng = f"{L}{FIRST_DATA_ROW}:{L}{last}"
            for r in range(FIRST_DATA_ROW, last + 1):
                cell = sh.cell(row=r, column=i)
                cell.alignment = Alignment(vertical="top", wrap_text=kind == "text",
                                           horizontal="center" if kind in ("bool", "int?", "float") else None)
                if kind == "date":
                    cell.number_format = "DD/MM/YYYY"
                elif kind == "datetime":
                    cell.number_format = "DD/MM/YYYY HH:MM"
            if kind in LISTS:
                _add_list_validation(sh, rng, kind, list_ranges, strict=False)
            if kind in ("risque", "contact_risk", "ipc", "priority"):
                _risk_formatting(sh, rng)
        # mise en forme explicite (identique dans Excel, LibreOffice et Google Sheets)
        band = PatternFill("solid", fgColor="F3F6FA")
        for i in range(1, len(cols) + 1):
            h = sh.cell(row=HEADER_ROW, column=i)
            h.font, h.fill, h.border = Font(bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY), BOX
            for r in range(FIRST_DATA_ROW, last + 1):
                cell = sh.cell(row=r, column=i)
                cell.border = BOX
                if (r - FIRST_DATA_ROW) % 2 == 1:
                    cell.fill = band
        sh.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(cols))}{last}"
        sh.freeze_panes = f"A{FIRST_DATA_ROW}"
        _print_setup(sh, HEADER_ROW)

    # ---- aperçu ------------------------------------------------------------------------
    if preview_png and Path(preview_png).exists():
        from openpyxl.drawing.image import Image as XLImage
        ap = wb.create_sheet("Aperçu")
        ap.sheet_view.showGridLines = False
        ap["A1"] = "Aperçu du mapping généré par l'outil (cette feuille est ignorée à l'import)"
        ap["A1"].font = Font(bold=True, color=NAVY, size=12)
        img = XLImage(str(preview_png))
        img.width, img.height = 1280, 720
        ap.add_image(img, "A3")

    for name in ("Listes", "Ref_ZS", "Ref_AS"):
        wb[name].sheet_state = "hidden"
    wb.active = 0
    return wb


def _print_setup(ws, header_row: int, landscape: bool = True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4


def _readme(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 90
    ws.merge_cells("B1:C1")
    ws["B1"] = "MAPPING ÉPI — Modèle Excel de saisie du parcours d'un cas"
    ws["B1"].font = Font(bold=True, size=16, color=WHITE)
    for col in "ABC":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 34
    rows = [
        ("COMMENT L'UTILISER", None),
        ("1. Remplir", "Compléter la feuille « Cas », puis une ligne par étape dans « Chronologie », et les feuilles "
                       "Lieux, Contacts, PCI, Laboratoire, Actions."),
        ("2. Importer", "Dans l'outil : « Importer une fiche » (version web) ou « Excel ▸ Importer une fiche Excel » "
                        "(application de bureau). Le mapping se génère automatiquement."),
        ("3. Télécharger", "Depuis l'outil : PNG, PDF, PowerPoint modifiable, SVG."),
        ("RÈGLES DE SAISIE", None),
        ("Inconnu", "Laisser la case VIDE (ou « ? » pour un nombre de contacts). Ne jamais mettre 0 quand on ne sait pas."),
        ("Dates", "JJ/MM/AAAA (ex. 18/08/2026). Laboratoire : JJ/MM/AAAA HH:MM (ex. 30/08/2026 21:00)."),
        ("Période", "Une étape sur plusieurs jours : « Date début » + « Date fin » (ex. 01/08 → 14/08)."),
        ("Déplacement", "« Lieu » = départ, « Lieu d'arrivée » = destination : la carte affiche « A → B »."),
        ("Oui / non", "Colonnes Contact corps, Liquides, Sans PCI, Sécurisé, Heure approx. : écrire « oui » ou laisser vide."),
        ("Risque", "Laisser vide pour que l'outil propose le risque selon ses règles (funérailles, contact avec le corps, "
                   "absence de PCI…). Remplir pour imposer une valeur."),
        ("Listes", "Province → Zone de santé → Aire de santé : listes officielles de la RDC (GRID3/HDX), filtrées en "
                   "cascade. Une aire absente de la liste peut être tapée librement."),
        ("Ne pas modifier", "Les noms des feuilles et les en-têtes de colonnes (ligne 3) : l'import s'appuie dessus. "
                            "On peut ajouter des lignes, pas des colonnes."),
        ("COULEURS DES NIVEAUX DE RISQUE", None),
    ]
    r = 3
    for k, v in rows:
        if v is None:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws.cell(row=r, column=2, value=k).font = Font(bold=True, color=NAVY, size=12)
            ws.cell(row=r, column=2).fill = SECTION_FILL
            ws.cell(row=r, column=3).fill = SECTION_FILL
            r += 1
            continue
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32
        r += 1
    for label, color, desc in [("Très élevé / Très haut / Non", theme.RED, "Contact avec le corps, funérailles, site non décontaminé"),
                               ("Élevé / Haut / Partiel", theme.ORANGE, "Soins, hospitalisation, rassemblement avec un cas symptomatique"),
                               ("Modéré / Faible / Contrôlé", theme.GREEN, "Exposition limitée ou circuit encadré"),
                               ("Inconnu / À investiguer / En cours", theme.BLUE, "Information à compléter")]:
        c = ws.cell(row=r, column=2, value=label)
        c.fill = PatternFill("solid", fgColor=color[1:])
        c.font = Font(bold=True, color=WHITE)
        ws.cell(row=r, column=3, value=desc)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value=f"Modèle généré le {datetime.now():%d/%m/%Y} — Mapping Épi").font = \
        Font(italic=True, size=9, color=MUTED)
    from config import credits
    for line in (credits.INSTITUTION, credits.EQUIPE, "Pour tout contact : " + ", ".join(credits.CONTACTS)):
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(row=r, column=2, value=line).font = Font(size=9, color=MUTED)
    _print_setup(ws, 1, landscape=False)


def save_workbook(case: Optional[Case], path: Path | str, preview_png: Optional[Path] = None) -> None:
    build_workbook(case, preview_png).save(str(path))
