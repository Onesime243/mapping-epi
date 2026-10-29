"""Stockage : SQLite (cas + historique des versions), JSON et Excel."""
from __future__ import annotations

import getpass
import sqlite3
import threading
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional, Tuple, get_args

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from config import theme
from core.models import Action, Case, ContactGroup, Event, IPCSite, LabSample, Place

ROOT = Path(__file__).resolve().parent.parent


def current_user() -> str:
    try:
        return getpass.getuser()
    except Exception:
        return "inconnu"


class Store:
    """Base locale partagée par l'application de bureau et le serveur web."""

    def __init__(self, path: Path | str | None = None, mirror_json: bool = True):
        if path is None:
            from core.paths import db_path
            path = db_path()
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.mirror_dir = (self.path.parent / "cas") if mirror_json else None
        self.lock = threading.RLock()
        self.db = sqlite3.connect(str(self.path), check_same_thread=False, timeout=15)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.executescript(
            """
            CREATE TABLE IF NOT EXISTS cases(
                id TEXT PRIMARY KEY, titre TEXT, zs TEXT, maladie TEXT,
                updated TEXT, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS versions(
                vid INTEGER PRIMARY KEY AUTOINCREMENT, case_id TEXT, ts TEXT,
                auteur TEXT, commentaire TEXT, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS users(
                username TEXT PRIMARY KEY, nom TEXT, pwd TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'agent',
                actif INTEGER NOT NULL DEFAULT 1, zs TEXT DEFAULT '', cree_le TEXT, derniere_connexion TEXT);
            """
        )
        cols = [r[1] for r in self.db.execute("PRAGMA table_info(cases)")]
        if "owner" not in cols:     # migration v1.0 → v1.1
            self.db.execute("ALTER TABLE cases ADD COLUMN owner TEXT DEFAULT ''")
        self.db.commit()

    # -- cas -------------------------------------------------------------
    def save(self, case: Case, version: bool = False, auteur: str = "", commentaire: str = "",
             owner: Optional[str] = None) -> None:
        case.modifie_le = datetime.now().replace(microsecond=0)
        case.modifie_par = auteur or case.investigateur or current_user()
        data = case.model_dump_json()
        with self.lock:
            row = self.db.execute("SELECT owner FROM cases WHERE id=?", (case.id,)).fetchone()
            own = owner if owner is not None else (row[0] if row else "")
            self.db.execute(
                "INSERT INTO cases(id,titre,zs,maladie,updated,data,owner) VALUES(?,?,?,?,?,?,?) "
                "ON CONFLICT(id) DO UPDATE SET titre=excluded.titre, zs=excluded.zs, "
                "maladie=excluded.maladie, updated=excluded.updated, data=excluded.data, owner=excluded.owner",
                (case.id, case.titre_court(), case.zs, case.maladie, case.modifie_le.isoformat(), data, own or ""),
            )
            if version:
                self.db.execute(
                    "INSERT INTO versions(case_id,ts,auteur,commentaire,data) VALUES(?,?,?,?,?)",
                    (case.id, case.modifie_le.isoformat(), case.modifie_par, commentaire, data),
                )
            self.db.commit()
        if self.mirror_dir:
            try:
                self.mirror_dir.mkdir(exist_ok=True)
                (self.mirror_dir / f"{case.id}.json").write_text(case.model_dump_json(indent=2), encoding="utf-8")
            except OSError:
                pass

    def load(self, case_id: str) -> Optional[Case]:
        row = self.db.execute("SELECT data FROM cases WHERE id=?", (case_id,)).fetchone()
        return Case.model_validate_json(row[0]) if row else None

    def owner_of(self, case_id: str) -> Optional[str]:
        row = self.db.execute("SELECT owner FROM cases WHERE id=?", (case_id,)).fetchone()
        return row[0] if row else None

    def list(self, search: str = "", owner: Optional[str] = None) -> List[Tuple[str, str, str]]:
        q = "SELECT id, titre, updated FROM cases WHERE 1=1"
        args: list = []
        if search:
            q += " AND (titre LIKE ? OR data LIKE ?)"
            args += [f"%{search}%", f"%{search}%"]
        if owner is not None:
            q += " AND owner=?"
            args.append(owner)
        return self.db.execute(q + " ORDER BY updated DESC", args).fetchall()

    def list_full(self, owner: Optional[str] = None):
        q = "SELECT id, titre, updated, owner FROM cases"
        args: tuple = ()
        if owner is not None:
            q += " WHERE owner=?"
            args = (owner,)
        return self.db.execute(q + " ORDER BY updated DESC", args).fetchall()

    def all_cases(self, owner: Optional[str] = None) -> List[Case]:
        if owner is None:
            rows = self.db.execute("SELECT data FROM cases")
        else:
            rows = self.db.execute("SELECT data FROM cases WHERE owner=?", (owner,))
        return [Case.model_validate_json(r[0]) for r in rows]

    def delete(self, case_id: str) -> None:
        # la fiche reste récupérable dans l'historique des versions
        with self.lock:
            self.db.execute("DELETE FROM cases WHERE id=?", (case_id,))
            self.db.commit()
        if self.mirror_dir:
            f = self.mirror_dir / f"{case_id}.json"
            if f.exists():
                f.rename(f.with_suffix(".supprime.json"))

    def versions(self, case_id: str):
        return self.db.execute(
            "SELECT vid, ts, auteur, commentaire FROM versions WHERE case_id=? ORDER BY vid DESC",
            (case_id,)).fetchall()

    def load_version(self, vid: int) -> Case:
        row = self.db.execute("SELECT data FROM versions WHERE vid=?", (vid,)).fetchone()
        return Case.model_validate_json(row[0])

    def close(self):
        self.db.close()


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------
def save_json(case: Case, path: Path | str) -> None:
    Path(path).write_text(case.model_dump_json(indent=2), encoding="utf-8")


def load_json(path: Path | str) -> Case:
    return Case.model_validate_json(Path(path).read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# Comparaison de versions
# ---------------------------------------------------------------------------
def diff_cases(old: Case, new: Case) -> List[str]:
    out = []
    a, b = old.model_dump(mode="json"), new.model_dump(mode="json")
    skip = {"modifie_le", "modifie_par"}
    for k in b:
        if k in skip:
            continue
        if isinstance(b[k], list):
            la, lb = a.get(k) or [], b[k]
            if la != lb:
                ida = {x.get("id"): x for x in la if isinstance(x, dict)}
                idb = {x.get("id"): x for x in lb if isinstance(x, dict)}
                added = len(set(idb) - set(ida))
                removed = len(set(ida) - set(idb))
                changed = sum(1 for i in set(ida) & set(idb) if ida[i] != idb[i])
                if ida or idb:
                    out.append(f"{k} : {added} ajout(s), {removed} suppression(s), {changed} modification(s)")
                else:
                    out.append(f"{k} : {la} → {lb}")
        elif a.get(k) != b[k]:
            out.append(f"{k} : « {a.get(k)} » → « {b[k]} »")
    return out or ["Aucune différence."]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
_CASE_FIELDS = [
    ("rang", "Rang du cas"), ("maladie", "Maladie"), ("sexe", "Sexe"), ("age", "Âge"),
    ("age_unite", "Unité âge"), ("profession", "Profession"), ("province", "Province"),
    ("zs", "Zone de santé"), ("aire_sante", "Aire de santé"), ("village", "Village / quartier"),
    ("lat", "Latitude"), ("lon", "Longitude"), ("statut", "Statut"),
    ("date_confirmation", "Date confirmation"), ("issue", "Issue"), ("investigateur", "Investigateur"),
    ("date_investigation", "Date investigation"), ("debut_symptomes", "Début symptômes"),
    ("symptomes", "Symptômes (séparés par ;)"), ("symptomes_autres", "Autres symptômes"),
    ("date_deces", "Date décès"), ("risque_global", "Risque global (vide = auto)"),
]

SHEETS = {
    "Chronologie": (Event, [
        ("date_debut", "Date début"), ("date_fin", "Date fin"), ("heure", "Heure"),
        ("heure_approx", "Heure approx."), ("date_incertaine", "Date incertaine"), ("lieu", "Lieu"),
        ("lieu_arrivee", "Lieu d'arrivée"), ("type", "Type"), ("titre", "Titre"),
        ("description", "Description"), ("transport", "Transport"), ("personnes", "Personnes présentes"),
        ("contact_corps", "Contact corps"), ("liquides", "Liquides biologiques"), ("sans_pci", "Sans PCI"),
        ("securise", "Sécurisé"), ("risque", "Risque (vide = auto)")]),
    "Lieux": (Place, [
        ("nom", "Lieu"), ("type", "Type"), ("population", "Population exposée"), ("nombre", "Nombre estimé"),
        ("risque", "Risque"), ("lat", "Latitude"), ("lon", "Longitude")]),
    "Contacts": (ContactGroup, [
        ("categorie", "Catégorie"), ("lien", "Lien avec le cas"), ("nombre", "Nombre (? si inconnu)"),
        ("libelle_nombre", "Libellé"), ("suivis", "Suivis"), ("risque", "Risque"), ("lieu", "Lieu de contact"),
        ("dernier_contact", "Dernier contact")]),
    "PCI": (IPCSite, [
        ("site", "Site"), ("statut", "Statut"), ("date", "Date"), ("equipe", "Équipe"), ("produit", "Produit")]),
    "Laboratoire": (LabSample, [
        ("type", "Échantillon"), ("laboratoire", "Laboratoire"), ("prelevement", "Prélèvement"),
        ("envoi", "Envoi"), ("reception", "Réception"), ("resultat_le", "Résultat le"), ("resultat", "Résultat")]),
    "Actions": (Action, [
        ("texte", "Action"), ("priorite", "Priorité"), ("responsable", "Responsable"),
        ("echeance", "Échéance"), ("statut", "Statut")]),
}
# colonnes pré-remplies par défaut (listes) qui ne suffisent pas à faire une ligne
DEFAULTED = {"Chronologie": ("type",), "Contacts": ("risque",), "PCI": ("statut",), "Actions": ("priorite", "statut")}
_LIST_ATTR = {"Chronologie": "events", "Lieux": "places", "Contacts": "contacts", "PCI": "ipc",
              "Laboratoire": "labs", "Actions": "actions"}
_CHOICES = {
    ("Chronologie", "type"): theme.EVENT_TYPE_CHOICES, ("Chronologie", "risque"): theme.RISK_CHOICES,
    ("Lieux", "risque"): theme.RISK_CHOICES, ("Lieux", "type"): theme.PLACE_TYPES,
    ("Contacts", "risque"): theme.CONTACT_RISK_CHOICES, ("PCI", "statut"): theme.IPC_CHOICES,
    ("Actions", "priorite"): theme.ACTION_PRIORITY_CHOICES, ("Actions", "statut"): theme.ACTION_STATUS_CHOICES,
}

_HEAD_FILL = PatternFill("solid", fgColor=theme.NAVY[1:])
_HEAD_FONT = Font(bold=True, color="FFFFFF")


def _to_cell(v):
    if isinstance(v, bool):
        return "oui" if v else ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, list):
        return "; ".join(v)
    return v


def parse_date(v) -> Optional[date]:
    if v in (None, "", "?"):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Date invalide : « {s} » (format attendu JJ/MM/AAAA)")


def parse_datetime(v) -> Optional[datetime]:
    if v in (None, "", "?"):
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip().replace("h", ":").rstrip(":")
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"Date/heure invalide : « {v} » (format attendu JJ/MM/AAAA HH:MM)")


def _field_type(model, key):
    ann = model.model_fields[key].annotation
    args = [a for a in get_args(ann) if a is not type(None)]
    optional = bool(get_args(ann)) and type(None) in get_args(ann)
    return (args[0] if args else ann), optional


def _from_cell(model, key, v):
    t, optional = _field_type(model, key)
    if v is None or (isinstance(v, str) and v.strip() == ""):
        if t is bool:
            return False
        return None if optional else ""
    if t is bool:
        return str(v).strip().lower() in ("oui", "o", "x", "1", "true", "vrai", "yes")
    if t is datetime:
        return parse_datetime(v)
    if t is date:
        return parse_date(v)
    if t is int:
        if str(v).strip() == "?":
            return None
        return int(float(v))
    if t is float:
        return float(str(v).replace(",", "."))
    return str(v).strip()


def export_excel(case: Case, path: Path | str, preview_png: Optional[Path] = None) -> None:
    """Classeur mis en forme (modèle de saisie), réimportable tel quel : voir core/excel_model.py."""
    from core.excel_model import save_workbook
    save_workbook(case, path, preview_png)


def import_excel(path: Path | str, base: Optional[Case] = None) -> Case:
    wb = load_workbook(str(path), data_only=True)
    data = base.model_copy(deep=True) if base else Case()
    errors: List[str] = []
    if "Cas" in wb.sheetnames:
        labels = {label: key for key, label in _CASE_FIELDS}
        for row in wb["Cas"].iter_rows(min_row=1, values_only=True):
            if not row or len(row) < 2 or row[0] not in labels:
                continue
            key, v = labels[row[0]], row[1]
            try:
                if key == "symptomes":
                    setattr(data, key, [s.strip() for s in str(v or "").split(";") if s.strip()])
                else:
                    setattr(data, key, _from_cell(Case, key, v))
            except ValueError as exc:
                errors.append(f"Cas / {row[0]} : {exc}")
    for sheet, (model, cols) in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        # la ligne d'en-têtes est la première (parmi les 6 premières) qui contient le libellé de la 1re colonne
        hrow = 0
        for i, r in enumerate(rows[:6]):
            if cols[0][1] in [str(h or "").strip() for h in r]:
                hrow = i
                break
        header = [str(h or "").strip() for h in rows[hrow]]
        idx = {key: header.index(label) for key, label in cols if label in header}
        items = []
        for n, r in enumerate(rows[hrow + 1:], start=hrow + 2):
            if not r or all(v in (None, "") for v in r):
                continue
            kw = {}
            for key, i in idx.items():
                try:
                    kw[key] = _from_cell(model, key, r[i] if i < len(r) else None)
                except ValueError as exc:
                    errors.append(f"{sheet} ligne {n} : {exc}")
            kw = {k: v for k, v in kw.items() if v is not None or _field_type(model, k)[1]}
            first_key = cols[0][0]
            if not str(kw.get(first_key) or "").strip() and not any(
                    v not in (None, "", False) for k, v in kw.items() if k != first_key and k not in DEFAULTED.get(sheet, ())):
                continue                       # ligne sans contenu utile
            items.append(model(**kw))
        setattr(data, _LIST_ATTR[sheet], items)
    if errors:
        raise ExcelImportError(data, errors)
    return data


class ExcelImportError(Exception):
    def __init__(self, case: Case, errors: List[str]):
        super().__init__("\n".join(errors))
        self.case = case
        self.errors = errors


def import_contact_list(path: Path | str) -> List[ContactGroup]:
    """Liste nominative (colonnes : Nom, Catégorie, Lieu, Dernier contact, Suivi [oui/non], Risque)
    agrégée par catégorie."""
    wb = load_workbook(str(path), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    ic, il, id_, isv, ir = (col("catégorie", "categorie"), col("lieu", "lieu de contact"),
                            col("dernier contact", "date dernier contact"), col("suivi", "suivis"), col("risque"))
    groups = defaultdict(lambda: {"n": 0, "s": 0, "lieu": "", "d": None, "r": "Haut"})
    for r in rows[1:]:
        if not r or all(v in (None, "") for v in r):
            continue
        cat = str(r[ic] if ic is not None and r[ic] else "Non classé")
        g = groups[cat]
        g["n"] += 1
        if isv is not None and str(r[isv] or "").lower() in ("oui", "o", "1", "x", "yes"):
            g["s"] += 1
        if il is not None and r[il]:
            g["lieu"] = str(r[il])
        if id_ is not None and r[id_]:
            try:
                d = parse_date(r[id_])
                g["d"] = max(g["d"], d) if g["d"] else d
            except ValueError:
                pass
        if ir is not None and r[ir] and str(r[ir]) in theme.CONTACT_RISK_CHOICES:
            if str(r[ir]) == "Très haut":
                g["r"] = "Très haut"
    return [ContactGroup(categorie=k, nombre=v["n"], suivis=v["s"], lieu=v["lieu"],
                         dernier_contact=v["d"], risque=v["r"], libelle_nombre="listés")
            for k, v in groups.items()]


def export_excel_template(path: Path | str) -> None:
    export_excel(Case(), path)
